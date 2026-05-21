"""Bundle generator: Excel workbook + Sponsor email from validated YAML blocks."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from pman.validator import ProjectValidator


def _val(v):
    return ", ".join(v) if isinstance(v, list) else str(v)


class BundleGenerator:
    """Generates a project artifact bundle (Excel + email) from plan content."""

    def __init__(self, project_name: str, content: str, output_dir: Path):
        self.project_name = project_name
        self.content = content
        self.output_dir = output_dir
        self.validator = ProjectValidator()
        self.blocks = self.validator.parse_yaml_blocks(content)

    def generate_excel(self) -> Path:
        """Create project_artifacts.xlsx with WBS, RACI, Risks, Budget sheets."""
        wb = Workbook()
        if wb.active:
            wb.remove(wb.active)

        self._add_wbs_sheet(wb)
        self._add_raci_sheet(wb)
        self._add_risks_sheet(wb)
        self._add_budget_sheet(wb)

        path = self.output_dir / "project_artifacts.xlsx"
        wb.save(path)
        return path

    def _add_wbs_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("WBS")
        ws.append(["ID", "Name", "Duration (days)", "Dependencies"])
        ws[1][0].font = Font(bold=True)
        for task in self.blocks.get("wbs", {}).get("wbs", []):
            ws.append([
                task.get("id", ""),
                task.get("name", ""),
                task.get("duration_days", ""),
                str(task.get("dependencies", [])),
            ])

    def _add_raci_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("RACI")
        ws.append(["Role", "Responsible", "Accountable", "Consulted", "Informed"])
        ws[1][0].font = Font(bold=True)
        for row in self.blocks.get("raci", {}).get("raci_matrix", []):
            ws.append([
                _val(row.get("role", "")),
                _val(row.get("responsible", "")),
                _val(row.get("accountable", "")),
                _val(row.get("consulted", "")),
                _val(row.get("informed", "")),
            ])

    def _add_risks_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Risks")
        ws.append(["Description", "Probability", "Impact", "Mitigation", "Owner"])
        ws[1][0].font = Font(bold=True)
        for risk in self.blocks.get("risks", {}).get("risks", []):
            ws.append([
                risk.get("description", ""),
                risk.get("probability", ""),
                risk.get("impact", ""),
                risk.get("mitigation", ""),
                risk.get("owner", ""),
            ])

    def _add_budget_sheet(self, wb: Workbook) -> None:
        ws = wb.create_sheet("Budget")
        ws.append(["Item", "Estimated Cost", "Actual Cost"])
        ws[1][0].font = Font(bold=True)
        budget = self.blocks.get("budget", {}).get("budget", {})
        for item in budget.get("items", []):
            ws.append([
                item.get("name", ""),
                item.get("estimated_cost", 0),
                item.get("actual_cost", 0),
            ])
        ws.append(["", "Total:", budget.get("total", 0)])

    def generate_email(self) -> Path:
        """Create sponsor_email.txt with key project metrics."""
        metrics = self.to_dict()
        template = (
            f"Subject: Project Update — {self.project_name}\n\n"
            "Dear Sponsor,\n\n"
            f"Please find attached the project plan artifacts for "
            f'"{self.project_name}".\n\n'
            "Key Metrics:\n"
            f"- Total Budget: {metrics['budget']['total']}\n"
            f"- Number of WBS Tasks: {len(metrics['wbs'])}\n"
            f"- Top Risk: "
            f"{metrics['risks'][0]['description'] if metrics['risks'] else 'N/A'} "
            f"(Probability: "
            f"{metrics['risks'][0]['probability'] if metrics['risks'] else 'N/A'})\n\n"
            "The full project plan, WBS, RACI matrix, risk register, "
            "and budget breakdown are available in the attached Excel workbook.\n\n"
            "Best regards,\nProject Manager\n"
        )

        path = self.output_dir / "sponsor_email.txt"
        path.write_text(template, encoding="utf-8")
        return path

    def to_dict(self) -> dict:
        """Return bundle data as a structured dict for JSON export."""
        wbs = self.blocks.get("wbs", {}).get("wbs", [])
        raci = self.blocks.get("raci", {}).get("raci_matrix", [])
        risks = self.blocks.get("risks", {}).get("risks", [])
        budget = self.blocks.get("budget", {}).get("budget", {})
        return {
            "project_name": self.project_name,
            "wbs": [
                {
                    "id": t.get("id", ""),
                    "name": t.get("name", ""),
                    "duration_days": t.get("duration_days"),
                    "dependencies": t.get("dependencies", []),
                }
                for t in wbs
            ],
            "raci": [
                {
                    "role": _val(r.get("role", "")),
                    "responsible": _val(r.get("responsible", "")),
                    "accountable": _val(r.get("accountable", "")),
                    "consulted": _val(r.get("consulted", "")),
                    "informed": _val(r.get("informed", "")),
                }
                for r in raci
            ],
            "risks": [
                {
                    "description": r.get("description", ""),
                    "probability": r.get("probability", ""),
                    "impact": r.get("impact", ""),
                    "mitigation": r.get("mitigation", ""),
                    "owner": r.get("owner", ""),
                }
                for r in risks
            ],
            "budget": {
                "items": [
                    {
                        "name": i.get("name", ""),
                        "estimated_cost": i.get("estimated_cost", 0),
                        "actual_cost": i.get("actual_cost", 0),
                    }
                    for i in budget.get("items", [])
                ],
                "total": budget.get("total", 0),
            },
        }
